"""Writes the content of the deck's companion workbook (openpyxl 3.1.5, on the box).

    python make_workbook.py OUT.xlsx

The same numbers as the deck's charts, with native bar, line, pie and combo charts, plus what a small
real workbook has: a merged title, number formats, formulas, a drop-down. The fixtures are what real
apps make of it (LibreOffice Calc, GridCraft, Excel); this file only writes the content.
"""
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "Round trips"
ws.merge_cells("A1:D1")
ws["A1"] = "A week of round trips"
ws["A1"].font = Font(bold=True, size=14)
ws["A1"].alignment = Alignment(horizontal="center")
ws.append(["Day", "Round trips", "Losses found", "Share found"])
for c in ws[2]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="97DBC2")
for day, trips, losses in [("Mon", 12, 5), ("Tue", 18, 3), ("Wed", 15, 4), ("Thu", 22, 2), ("Fri", 27, 1)]:
    r = ws.max_row + 1
    ws.append([day, trips, losses, f"=C{r}/B{r}"])
    ws[f"D{r}"].number_format = "0.0%"
ws.append(["Total", "=SUM(B3:B7)", "=SUM(C3:C7)", "=C8/B8"])
ws["D8"].number_format = "0.0%"
for c in ws[8]:
    c.font = Font(bold=True)
ws["F2"] = "Status"
ws["F3"] = "open"
dv = DataValidation(type="list", formula1='"open,fixed,listed"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("F3:F7")
ws.column_dimensions["A"].width = 12

cats = Reference(ws, min_col=1, min_row=3, max_row=7)
bar = BarChart()
bar.title = "Round trips and losses"
bar.add_data(Reference(ws, min_col=2, max_col=3, min_row=2, max_row=7), titles_from_data=True)
bar.set_categories(cats)
ws.add_chart(bar, "H2")

line = LineChart()
line.title = "The same week, as lines"
line.add_data(Reference(ws, min_col=2, max_col=3, min_row=2, max_row=7), titles_from_data=True)
line.set_categories(cats)
ws.add_chart(line, "H18")

# A combo: bars for round trips, a line for losses, on the same axes.
combo = BarChart()
combo.title = "Round trips (bars) and losses (line)"
combo.add_data(Reference(ws, min_col=2, min_row=2, max_row=7), titles_from_data=True)
combo.set_categories(cats)
over = LineChart()
over.add_data(Reference(ws, min_col=3, min_row=2, max_row=7), titles_from_data=True)
combo += over
ws.add_chart(combo, "H34")

fs = wb.create_sheet("Formats")
fs.append(["Format", "Files checked"])
for row in [("Word", 40), ("Excel", 35), ("PowerPoint", 25)]:
    fs.append(row)
fs.append(["Total", "=SUM(B2:B4)"])
pie = PieChart()
pie.title = "Files checked"
pie.add_data(Reference(fs, min_col=2, min_row=1, max_row=4), titles_from_data=True)
pie.set_categories(Reference(fs, min_col=1, min_row=2, max_row=4))
fs.add_chart(pie, "D2")

wb.properties.creator = "Milleleaf"
wb.properties.title = "Milleleaf charts"
wb.save(sys.argv[1])
print(sys.argv[1])
