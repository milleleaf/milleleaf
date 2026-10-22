"""A one-slide deck holding only a bar chart, for LibreOffice to render as the deck's chart picture.
    python chart_slide.py OUT.pptx"""
import sys
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(8), Inches(5)
s = prs.slides.add_slide(prs.slide_layouts[6])
d = CategoryChartData()
d.categories = ["Mon", "Tue", "Wed", "Thu", "Fri"]
d.add_series("Round trips", (12, 18, 15, 22, 27))
c = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.3), Inches(0.3), Inches(7.4), Inches(4.4), d).chart
c.chart_title.text_frame.text = "Round trips this week"
prs.save(sys.argv[1])
